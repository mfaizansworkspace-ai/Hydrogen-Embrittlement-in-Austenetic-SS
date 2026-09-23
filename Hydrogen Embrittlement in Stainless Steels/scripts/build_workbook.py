import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
OUT="/mnt/user-data/outputs/HE-Austenite_v0.1.xlsx"
d=json.load(open("data/clean.json")); rows=d["rows"]; SRC=d["sources"]
A=lambda **k: Font(name="Arial",**k); HDR=A(bold=True,color="FFFFFF"); FILL=PatternFill("solid",fgColor="1F3864")
wb=Workbook()
def head(ws,n):
    for c in range(1,n+1):
        x=ws.cell(1,c); x.font=HDR; x.fill=FILL
        x.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
# Read me
rm=wb.active; rm.title="Read me"
for i,(t,sz,b) in enumerate([("HE-Austenite v0.1",14,True),
 ("Hydrogen embrittlement of austenitic stainless steels in gaseous hydrogen: base metal and welds",11,False),("",10,False),
 ("Scope: gaseous hydrogen only. Tested in high-pressure hydrogen gas, or thermally precharged in hydrogen gas and then tested.",10,False),
 ("Electrochemically charged studies are excluded: cathodic fugacity cannot be converted to an equivalent gas pressure without assumptions.",10,False),("",10,False),
 ("Status: draft. The 'verified' column marks rows that were independently re-extracted and matched.",10,False),
 ("Derived values (nickel equivalent, relative ratio) are flagged in the 'derived' column and were not reported as such by the source.",10,False),
 ("Sources whose data appears only in figures are listed in the Extraction log sheet and are not in the Data sheet.",10,False),("",10,False),
 ("Compiled by Muhammad Faizan, September 2026. Data CC BY 4.0; code MIT.",10,False)],1):
    rm[f"A{i}"]=t; rm[f"A{i}"].font=A(size=sz,bold=b)
rm.column_dimensions["A"].width=130
# Data
ws=wb.create_sheet("Data"); cols=list(rows[0].keys())
ws.append([c.replace("_"," ") for c in cols])
for r in rows: ws.append([r[c] for c in cols])
thin=Side(style="thin",color="D9D9D9")
for row in ws.iter_rows(min_row=2):
    for c in row: c.font=A(size=10); c.border=Border(bottom=thin)
ci={c:i+1 for i,c in enumerate(cols)}; last=len(rows)+1
for key,w in dict(record_id=26,source_id=9,source_short=15,alloy=22,heat=8,material_family=27,product_form=12,
  is_weld=9,zone=13,weld_process=13,filler_metal=12,ferrite_number=13,hydrogen_exposure=18,h2_pressure_MPa=14,
  hydrogen_content_wppm=16,charging_time_h=13,test_temperature_K=14,test_temperature_C=14,strain_rate_per_s=14,
  test_type=11,property_name=20,reference_environment=16,property_reference=14,property_h2=12,relative_ratio=12,
  embrittlement_flag=18,ni_equivalent=12,derived=22,verified=9,notes=70).items():
    ws.column_dimensions[get_column_letter(ci[key])].width=w
rrL=get_column_letter(ci["relative_ratio"]); propL=get_column_letter(ci["property_name"])
weldL=get_column_letter(ci["is_weld"]); famL=get_column_letter(ci["material_family"])
flagL=get_column_letter(ci["embrittlement_flag"]); verL=get_column_letter(ci["verified"])
n=len(cols)
ws.cell(1,n+1,"RA ratio (base metal)"); ws.cell(1,n+2,"RA ratio (weld)")
for r in range(2,last+1):
    ws.cell(r,n+1,f'=IF(AND(${propL}{r}="reduction of area",${weldL}{r}="no"),${rrL}{r},"")').number_format="0.000"
    ws.cell(r,n+2,f'=IF(AND(${propL}{r}="reduction of area",${weldL}{r}="yes"),${rrL}{r},"")').number_format="0.000"
    for c in (n+1,n+2): ws.cell(r,c).font=A(size=10)
for c in (n+1,n+2): ws.column_dimensions[get_column_letter(c)].width=16
head(ws,n+2)
B=get_column_letter(n+1); W=get_column_letter(n+2)
# Summary
s=wb.create_sheet("Summary")
s["A1"]="HE-Austenite v0.1 — summary"; s["A1"].font=A(bold=True,size=14)
s["A2"]="Gaseous hydrogen only. All figures below are live formulas over the Data sheet."; s["A2"].font=A(size=10,italic=True)
put=lambda cell,v,**k: (s.__setitem__(cell,v), setattr(s[cell],"font",A(**k)))
put("A4","Counts",bold=True)
for i,(lab,f) in enumerate([("Total records",f"=COUNTA(Data!A2:A{last})"),
  ("Reduction-of-area rows",f"=COUNT(Data!{B}2:{B}{last})+COUNT(Data!{W}2:{W}{last})"),
  ("Weld rows",f'=COUNTIF(Data!{weldL}2:{weldL}{last},"yes")'),
  ("Base metal rows",f'=COUNTIF(Data!{weldL}2:{weldL}{last},"no")'),
  ("Rows independently re-extracted and matched",f'=COUNTIF(Data!{verL}2:{verL}{last},"yes")')],5):
    s[f"A{i}"]=lab; s[f"A{i}"].font=A(size=10); s[f"B{i}"]=f
put("A11","Median RRA",bold=True)
s["A12"]="Base metal"; s["B12"]=f"=MEDIAN(Data!{B}2:{B}{last})"; s["B12"].number_format="0.00"
s["A13"]="Welds";      s["B13"]=f"=MEDIAN(Data!{W}2:{W}{last})"; s["B13"].number_format="0.00"
put("A15","By material family",bold=True)
for c,t in zip("ABCD",["Family","RA rows","Mean RRA","Rows with significant loss (RRA < 0.80)"]):
    s[f"{c}16"]=t; s[f"{c}16"].font=A(bold=True)
fams=sorted({r["material_family"] for r in rows})
for i,fam in enumerate(fams,17):
    s[f"A{i}"]=fam; s[f"A{i}"].font=A(size=10)
    s[f"B{i}"]=f'=COUNTIFS(Data!${famL}$2:${famL}${last},A{i},Data!${propL}$2:${propL}${last},"reduction of area")'
    s[f"C{i}"]=f'=IFERROR(AVERAGEIFS(Data!${rrL}$2:${rrL}${last},Data!${famL}$2:${famL}${last},A{i},Data!${propL}$2:${propL}${last},"reduction of area"),"")'
    s[f"D{i}"]=f'=COUNTIFS(Data!${famL}$2:${famL}${last},A{i},Data!${propL}$2:${propL}${last},"reduction of area",Data!${flagL}$2:${flagL}${last},"significant loss")'
    s[f"C{i}"].number_format="0.00"
    for c in "BCD": s[f"{c}{i}"].font=A(size=10)
s.column_dimensions["A"].width=42
for c,w in zip("BCD",(12,12,34)): s.column_dimensions[c].width=w
# Data dictionary
dd=wb.create_sheet("Data dictionary"); dd.append(["Column","Meaning","Units / values"])
for a,b,c in [("record_id","Unique id: source, table, row",""),("source_id","Key into the Sources sheet",""),
 ("source_short","Author and year",""),("alloy","Alloy designation, standardised",""),("heat","Heat or cast id where reported",""),
 ("material_family","Grouping used for analysis",""),("product_form","plate, bar, sheet, tube, weld",""),
 ("is_weld","Gauge section contains weld metal","yes / no"),("zone","Part of the joint tested","base metal / weld metal / HAZ / whole joint"),
 ("weld_process","Welding process","GTAW, GMAW, PAW, laser"),("filler_metal","Filler wire",""),
 ("ferrite_number","Ferrite number where reported","FN"),("hydrogen_exposure","How hydrogen was introduced","gaseous in-situ / gaseous precharge"),
 ("h2_pressure_MPa","Hydrogen gas pressure","MPa"),("hydrogen_content_wppm","Hydrogen content","wt ppm"),
 ("charging_time_h","Precharging duration","hours"),("test_temperature_K","Test temperature","K"),
 ("test_temperature_C","Test temperature","deg C"),("strain_rate_per_s","Nominal strain rate","1/s"),
 ("test_type","Test method","tensile / SSRT"),("property_name","Property measured",""),
 ("reference_environment","What hydrogen is compared against","air / helium / nitrogen / argon / uncharged"),
 ("property_reference","Value in the reference environment","% or MPa"),("property_h2","Value in or after hydrogen","% or MPa"),
 ("relative_ratio","property_h2 / property_reference","1.0 = no loss"),
 ("embrittlement_flag","Banding of relative_ratio","no measurable loss >=0.95 / moderate 0.80-0.95 / significant <0.80"),
 ("ni_equivalent","Nickel equivalent (Hirayama) where composition was reported","wt %"),
 ("derived","Fields computed rather than reported",""),("verified","Independently re-extracted and matched","yes / no"),
 ("notes","Source table, conditions, caveats","")]: dd.append([a,b,c])
for row in dd.iter_rows(min_row=2):
    for c in row: c.font=A(size=10); c.alignment=Alignment(vertical="top",wrap_text=True)
for c,w in zip("ABC",(24,60,46)): dd.column_dimensions[c].width=w
head(dd,3)
# Sources
sh=wb.create_sheet("Sources"); sh.append(["Source id","Citation","Link","Access"])
for k,v in sorted(SRC.items(),key=lambda kv: kv[1]["id"]): sh.append([v["id"],v["cite"],v["link"],v["access"]])
for row in sh.iter_rows(min_row=2):
    for c in row: c.font=A(size=10); c.alignment=Alignment(vertical="top",wrap_text=True)
for c,w in zip("ABCD",(11,88,44,12)): sh.column_dimensions[c].width=w
head(sh,4)
# Extraction log
el=wb.create_sheet("Extraction log"); el.append(["Source","Status","Reason"])
for a,b,c in [("Caskey DP-1643","extracted","Tables 3, 6, 10, 11; compositions from Table A-1. Remaining data is in figures."),
 ("SAND2012-7321 chapters 2101, 2103, 2104, 2201, 2202, 1600","extracted","Smooth and notched tensile tables plus weld tables."),
 ("Balch 2015","extracted","Table 3, all six weld positions, both orientations."),
 ("Nakamura 2018","extracted","Base metal and 317L weld metal, as-welded and solution-treated."),
 ("Michler 2009","extracted","Table 2, two heats of 316L."),("Younes 2013","extracted","Table 4, RA loss at three hydrogen contents."),
 ("Iyer 1989","extracted","Table 1; charged in Ar-3%H2 gas at 1050 C, so in scope."),
 ("Fukunaga 2024","extracted","316CW and A286 values given in the text."),
 ("Matsuoka 2017 (SSP)","partly extracted","Only the Type 304 room-temperature RRA (0.37) is numeric; the RRA vs Nieq data is figure-only."),
 ("Yamabe 2017","not extracted","RRA appears only in figures. Ten steels with Nieq already computed; the best candidate for digitisation."),
 ("Zhang 2013","not extracted","RRA vs prestrain and Nieq is figure-only."),
 ("Hirata 2015","not extracted","RRA vs delta-ferrite ratio is figure-only. Would be the key gaseous delta-ferrite dataset."),
 ("Neuharth 2015","not extracted","Burst tests with wall-thickness reduction; controls come from the literature rather than paired specimens."),
 ("Fu 2020, Zhou 2022/2024, Yun 2023, Choi 2024, Xu 2021, Astafurova 2023, Maritsa 2026, Sabzi 2025","excluded","Electrochemical or cathodic charging, outside the gaseous-hydrogen scope.")]:
    el.append([a,b,c])
for row in el.iter_rows(min_row=2):
    for c in row: c.font=A(size=10); c.alignment=Alignment(vertical="top",wrap_text=True)
for c,w in zip("ABC",(58,18,86)): el.column_dimensions[c].width=w
head(el,3)
wb.save(OUT); print("saved",OUT)
